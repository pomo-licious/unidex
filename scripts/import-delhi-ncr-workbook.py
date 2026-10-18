#!/usr/bin/env python3
"""Read the audited workbook without modifying it; regenerate the website catalogue.

Usage: python scripts/import-delhi-ncr-workbook.py /absolute/path/to/workbook.xlsx
Requires openpyxl for read-only XLSX extraction. Node is used for normalization.
"""
import argparse
import datetime
import gzip
import hashlib
import json
from pathlib import Path
import subprocess

from openpyxl import load_workbook


def json_value(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return value


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--version', default='7')
    parser.add_argument('--audit-date', default='2026-10-18')
    args = parser.parse_args()
    root = Path(__file__).resolve().parent.parent
    workbook_bytes = args.workbook.read_bytes()
    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    metadata = {
        'sourceWorkbook': args.workbook.name,
        'sourceLibraryId': 'libfile_cd4f5961026c8191af6ad127e6240c5e',
        'sourceVersion': args.version,
        'sourceAuditDate': args.audit_date,
        'sha256': hashlib.sha256(workbook_bytes).hexdigest(),
        'valueMode': 'cached values; original formulas retained in the unchanged workbook',
    }
    current = {'metadata': metadata, 'sheets': {}}
    historical = {'metadata': metadata, 'sheets': {}}
    for sheet in workbook:
        rows = [
            {'row': index, 'values': [json_value(value) for value in values]}
            for index, values in enumerate(sheet.iter_rows(values_only=True), 1)
            if any(value is not None for value in values)
        ]
        target = historical if sheet.title in ['Source Tracker', 'Change Log'] else current
        target['sheets'][sheet.title] = rows
    source_path = root / 'scripts/data/delhi-ncr-workbook-v7.json.gz'
    source_path.parent.mkdir(parents=True, exist_ok=True)
    source_path.write_bytes(gzip.compress(
        (json.dumps(current, ensure_ascii=False, separators=(',', ':')) + '\n').encode(), mtime=0,
    ))
    archive_path = root / 'docs/data/delhi-ncr-workbook-v7-audit.json.gz'
    archive_path.parent.mkdir(parents=True, exist_ok=True)
    archive_path.write_bytes(gzip.compress(
        json.dumps(historical, ensure_ascii=False, separators=(',', ':')).encode(), mtime=0,
    ))
    subprocess.run(['node', 'scripts/generate-delhi-ncr-catalogue.mjs'], cwd=root, check=True)


if __name__ == '__main__':
    main()
